"""
Generate Microsoft PowerPoint presentation (.pptx) from presentation_slide_deck.md.
Smart Health Monitoring System — Digital Assignment 3 (DA3) Submission.
"""

import os
import re
import sys
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

ROOT_DIR = Path(__file__).resolve().parent.parent
SLIDE_DECK_MD = ROOT_DIR / "docs" / "submission" / "presentation_slide_deck.md"
OUTPUT_PPTX = ROOT_DIR / "docs" / "submission" / "Smart_Health_Monitoring_DA3_Presentation.pptx"


def parse_slide_deck(md_path: Path):
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Split by slide headers
    slide_chunks = re.split(r"\n##\s+(?:Slide\s+\d+:|\d+\.\s+Slide\s+\d+:)", content)
    slides_data = []

    for chunk in slide_chunks[1:]:  # Skip pre-slide preamble
        lines = chunk.strip().split("\n")
        header_line = lines[0].strip() if lines else "Slide"

        title_match = re.search(r"###\s+Slide Title\s*\n\*\*([^*]+)\*\*", chunk)
        slide_title = title_match.group(1).strip() if title_match else header_line

        message_match = re.search(r"###\s+Main Message\s*\n([^\n#]+)", chunk)
        main_message = message_match.group(1).strip() if message_match else ""

        bullets = []
        bullets_section = re.search(r"###\s+On-Slide Bullet Points\s*\n([\s\S]*?)(?=###|\Z)", chunk)
        if bullets_section:
            for b_line in bullets_section.group(1).strip().split("\n"):
                b_line = b_line.strip()
                if b_line.startswith("- ") or b_line.startswith("* "):
                    cleaned = re.sub(r"^\s*[-*]\s+", "", b_line)
                    cleaned = re.sub(r"\*\*([^*]+)\*\*", r"\1", cleaned)  # remove markdown bold
                    bullets.append(cleaned)

        notes = []
        notes_match = re.search(r"###\s+Speaker Notes\s*\n([\s\S]*?)(?=###|\Z)", chunk)
        if notes_match:
            for n_line in notes_match.group(1).strip().split("\n"):
                n_line = n_line.strip()
                if n_line.startswith(">"):
                    notes.append(n_line.lstrip("> ").strip(' "'))
        speaker_notes = "\n\n".join(notes)

        qna_match = re.search(r"###\s+Likely Examiner Question\s*\n([\s\S]*?)(?=\n---\n|\Z)", chunk)
        if qna_match:
            qna_text = qna_match.group(1).strip()
            speaker_notes += f"\n\n--- EXAMINER Q&A ---\n{qna_text}"

        slides_data.append({
            "header": header_line,
            "title": slide_title,
            "main_message": main_message,
            "bullets": bullets,
            "speaker_notes": speaker_notes,
        })

    return slides_data


def build_pptx(slides_data, output_path: Path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Color palette
    PRIMARY_COLOR = RGBColor(15, 76, 129)    # Deep Medical Blue
    ACCENT_COLOR = RGBColor(0, 150, 136)     # Teal Accent
    DARK_TEXT = RGBColor(33, 37, 41)         # Slate Charcoal
    MUTED_TEXT = RGBColor(108, 117, 125)     # Muted Grey

    for idx, sdata in enumerate(slides_data):
        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)

        # Header banner text box
        banner_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.2))
        tf_banner = banner_box.text_frame
        tf_banner.word_wrap = True

        p_num = tf_banner.paragraphs[0]
        p_num.text = f"SMART HEALTH MONITORING SYSTEM  |  SLIDE {idx + 1}"
        p_num.font.size = Pt(11)
        p_num.font.bold = True
        p_num.font.color.rgb = ACCENT_COLOR

        p_title = tf_banner.add_paragraph()
        p_title.text = sdata["title"]
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = PRIMARY_COLOR
        p_title.space_after = Pt(4)

        if sdata["main_message"]:
            p_sub = tf_banner.add_paragraph()
            p_sub.text = sdata["main_message"]
            p_sub.font.size = Pt(13)
            p_sub.font.italic = True
            p_sub.font.color.rgb = MUTED_TEXT

        # Content bullets box
        content_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.8))
        tf_content = content_box.text_frame
        tf_content.word_wrap = True

        for b_idx, bullet in enumerate(sdata["bullets"]):
            p = tf_content.add_paragraph() if b_idx > 0 else tf_content.paragraphs[0]
            p.text = f"•  {bullet}"
            p.font.size = Pt(15)
            p.font.color.rgb = DARK_TEXT
            p.space_after = Pt(10)

        # Add speaker notes
        if sdata["speaker_notes"]:
            notes_slide = slide.notes_slide
            tf_notes = notes_slide.notes_text_frame
            tf_notes.text = sdata["speaker_notes"]

    prs.save(str(output_path))
    print(f"[OK] Generated PowerPoint presentation with {len(slides_data)} slides at:")
    print(f"     {output_path}")


if __name__ == "__main__":
    data = parse_slide_deck(SLIDE_DECK_MD)
    build_pptx(data, OUTPUT_PPTX)
